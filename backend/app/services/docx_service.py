from __future__ import annotations

import io
import math
import re
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from defusedxml.ElementTree import fromstring
from PIL import Image, ImageDraw, ImageFont
from docx import Document as WordDocument
from docx.document import Document as WordDocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models import AgentRun, Block, DocumentAsset
from app.services.block_service import BlockError
from app.services.document_asset_service import DOCX_MIME_TYPE, DocumentAssetService, MAX_DOCX_BYTES
from app.services.drawio_diagrams import diagram_review_projection


_DOC_MARKER = re.compile(r"FFDOC\|([^|]+)\|(\d+)\|")
_BLOCK_MARKER = re.compile(r"FFBLOCK\|([^|]+)\|(\d+)\|")
_CAPTION_MARKER = "FFCAPTION|"
_ASSET_LINE = re.compile(r"^!\[([^\]]*)\]\(asset://([A-Za-z0-9_-]+)\)\s*$")
_TABLE_DIVIDER = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
_INLINE_TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|\[[^\]]+\]\([^)]+\))")


@dataclass
class ImportedImage:
    token: str
    file_name: str
    mime_type: str
    content: bytes


@dataclass
class ImportedBlock:
    block_id: str
    base_version: int
    title: str
    content: str
    images: list[ImportedImage]


def _set_font(run, name: str = "Microsoft YaHei", size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_hyperlink(paragraph, text: str, url: str):
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class DocxService:
    """Round-trip Block PRDs through a revision-bound DOCX exchange document."""

    def __init__(self, block_service, assets: DocumentAssetService):
        self.blocks = block_service
        self.assets = assets

    def export(self, document_id: str, actor: str) -> tuple[str, bytes, int]:
        with self.blocks.factory() as db:
            document = self.blocks.document(db, document_id, actor)
            artifact = self.assets.artifact_for_revision(db, document.id, document.revision)
            if artifact is not None:
                return artifact.file_name, self.assets.artifact_bytes(artifact), artifact.revision
            revision = document.revision
            block_rows = self.blocks.blocks(db, document)
            asset_rows = db.query(DocumentAsset).filter_by(document_id=document.id).all()
            content = self._build(document, block_rows, asset_rows)
            file_name = self._file_name(document.title)

        created: list[Path] = []
        try:
            with self.blocks.write() as db:
                current = self.blocks.document(db, document_id, actor)
                if current.revision != revision:
                    raise BlockError("导出期间文档已变化，请重试")
                artifact = self.assets.artifact_for_revision(db, document_id, revision)
                if artifact is None:
                    artifact = self.assets.store_artifact(
                        db,
                        document_id=document_id,
                        revision=revision,
                        actor=actor,
                        file_name=file_name,
                        content=content,
                        created_paths=created,
                    )
                return artifact.file_name, self.assets.artifact_bytes(artifact), revision
        except BaseException:
            for path in created:
                path.unlink(missing_ok=True)
            raise

    def import_document(
        self,
        document_id: str,
        actor: str,
        expected_revision: int,
        file_name: str,
        content: bytes,
    ) -> dict[str, object]:
        if not content or len(content) > MAX_DOCX_BYTES:
            raise BlockError("DOCX 文件必须小于 24 MB", 422)
        try:
            word = WordDocument(io.BytesIO(content))
        except Exception as error:
            raise BlockError("DOCX 文件无效或已损坏", 422) from error

        with self.blocks.factory() as db:
            document = self.blocks.document(db, document_id, actor)
            block_rows = self.blocks.blocks(db, document)
            known_assets = {
                asset.id: asset
                for asset in db.query(DocumentAsset).filter_by(document_id=document.id)
            }
            parsed = self._parse(word, document.id, block_rows, known_assets)

        created: list[Path] = []
        changed_ids: list[str] = []
        try:
            with self.blocks.write() as db:
                document = self.blocks.document(db, document_id, actor)
                if document.revision != expected_revision:
                    raise BlockError("文档已改变，请刷新后重新导入")
                if db.query(AgentRun).filter_by(document_id=document.id).filter(
                    AgentRun.status.in_(["queued", "running"])
                ).count():
                    raise BlockError("请先取消运行中的 AI 作业")
                for imported in parsed:
                    block = self.blocks.block(db, document, imported.block_id)
                    if block.version != imported.base_version:
                        raise BlockError(f"{block.title} 已有新版本，请重新导出后编辑")
                    imported_content = imported.content
                    for image in imported.images:
                        asset = self.assets.store_asset(
                            db,
                            document_id=document.id,
                            block_id=block.id,
                            actor=actor,
                            file_name=image.file_name,
                            mime_type=image.mime_type,
                            content=image.content,
                            metadata={"source": "docx_import"},
                            created_paths=created,
                        )
                        imported_content = imported_content.replace(
                            image.token, f"![{asset.file_name}](asset://{asset.id})"
                        )
                    imported_content = imported_content.strip() + ("\n" if imported_content.strip() else "")
                    values = {"title": imported.title, "content": imported_content}
                    if block.title != imported.title or block.content != imported_content:
                        self.blocks.save_in_transaction(
                            db,
                            document,
                            block,
                            imported.base_version,
                            values,
                            actor,
                            source="docx_import",
                        )
                        changed_ids.append(block.id)
                artifact = self.assets.store_artifact(
                    db,
                    document_id=document.id,
                    revision=document.revision,
                    actor=actor,
                    file_name=self._file_name(Path(file_name).stem or document.title),
                    content=content,
                    kind="docx_import",
                    created_paths=created,
                )
                result = {
                    "changed_block_ids": changed_ids,
                    "revision": document.revision,
                    "artifact_id": artifact.id,
                }
            return result
        except BaseException:
            for path in created:
                path.unlink(missing_ok=True)
            raise

    def _build(self, document, blocks: list[Block], assets: list[DocumentAsset]) -> bytes:
        word = WordDocument()
        word.core_properties.title = document.title
        word.core_properties.subject = "firstFlight Block PRD shared editing document"
        section = word.sections[0]
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)
        self._configure_styles(word)

        title = word.add_paragraph(style="Title")
        marker = title.add_run(f"FFDOC|{document.id}|{document.revision}|")
        marker.font.hidden = True
        visible = title.add_run(document.title)
        _set_font(visible, size=22)
        intro = word.add_paragraph("本文件用于 firstFlight PRD 的共享编辑。编辑后可回传系统，系统会按章节更新对应 Block 并保留版本历史。")
        intro.style = word.styles["Subtitle"]

        depth = self._depths(blocks)
        assets_by_block: dict[str | None, list[DocumentAsset]] = {}
        for asset in assets:
            assets_by_block.setdefault(asset.block_id, []).append(asset)
        for block in blocks:
            heading = word.add_paragraph(style=f"Heading {min(depth.get(block.id, 1), 6)}")
            marker = heading.add_run(f"FFBLOCK|{block.id}|{block.version}|")
            marker.font.hidden = True
            run = heading.add_run(block.title)
            _set_font(run, size=max(11, 17 - depth.get(block.id, 1)))
            referenced = self._append_markdown(
                word, block.content or "", {asset.id: asset for asset in assets_by_block.get(block.id, [])}
            )
            for asset in assets_by_block.get(block.id, []):
                if asset.id not in referenced:
                    self._append_asset(word, asset, asset.file_name)
            self._append_diagrams(word, block)

        if assets_by_block.get(None):
            word.add_heading("文档附件", level=1)
            for asset in assets_by_block[None]:
                self._append_asset(word, asset, asset.file_name)

        output = io.BytesIO()
        word.save(output)
        return output.getvalue()

    @staticmethod
    def _configure_styles(word: WordDocumentType) -> None:
        for style_name, size, bold in (
            ("Normal", 10.5, False),
            ("Title", 22, True),
            ("Subtitle", 10, False),
            ("Heading 1", 16, True),
            ("Heading 2", 14, True),
            ("Heading 3", 12.5, True),
            ("Heading 4", 11.5, True),
            ("Heading 5", 11, True),
            ("Heading 6", 10.5, True),
        ):
            style = word.styles[style_name]
            style.font.name = "Microsoft YaHei"
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
        normal = word.styles["Normal"].paragraph_format
        normal.space_after = Pt(6)
        normal.line_spacing = 1.2
        for index in range(1, 7):
            fmt = word.styles[f"Heading {index}"].paragraph_format
            fmt.space_before = Pt(12 if index <= 2 else 8)
            fmt.space_after = Pt(5)
            fmt.keep_with_next = True

    @staticmethod
    def _depths(blocks: list[Block]) -> dict[str, int]:
        by_id = {block.id: block for block in blocks}
        depths: dict[str, int] = {}
        for block in blocks:
            depth, parent, seen = 1, block.parent_id, {block.id}
            while parent and parent in by_id and parent not in seen:
                seen.add(parent)
                depth += 1
                parent = by_id[parent].parent_id
            depths[block.id] = min(depth, 6)
        return depths

    def _append_markdown(
        self, word: WordDocumentType, markdown: str, assets: dict[str, DocumentAsset]
    ) -> set[str]:
        referenced: set[str] = set()
        lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        index = 0
        fenced = False
        code_lines: list[str] = []
        while index < len(lines):
            line = lines[index]
            if line.strip().startswith(("```", "~~~")):
                if fenced:
                    paragraph = word.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.25)
                    run = paragraph.add_run("\n".join(code_lines))
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                    fenced, code_lines = False, []
                else:
                    fenced = True
                index += 1
                continue
            if fenced:
                code_lines.append(line)
                index += 1
                continue
            image = _ASSET_LINE.match(line.strip())
            if image and image.group(2) in assets:
                asset = assets[image.group(2)]
                self._append_asset(word, asset, image.group(1) or asset.file_name, referenced=True)
                referenced.add(asset.id)
                index += 1
                continue
            if index + 1 < len(lines) and "|" in line and _TABLE_DIVIDER.match(lines[index + 1]):
                rows = [self._split_table_row(line)]
                index += 2
                while index < len(lines) and "|" in lines[index] and lines[index].strip():
                    rows.append(self._split_table_row(lines[index]))
                    index += 1
                self._append_table(word, rows)
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                paragraph = word.add_heading(level=min(len(heading.group(1)) + 1, 6))
                self._append_inline(paragraph, heading.group(2))
            elif re.match(r"^\s*[-*+]\s+", line):
                paragraph = word.add_paragraph(style="List Bullet")
                self._append_inline(paragraph, re.sub(r"^\s*[-*+]\s+", "", line))
            elif re.match(r"^\s*\d+[.)]\s+", line):
                paragraph = word.add_paragraph(style="List Number")
                self._append_inline(paragraph, re.sub(r"^\s*\d+[.)]\s+", "", line))
            elif line.startswith(">"):
                paragraph = word.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                self._append_inline(paragraph, line[1:].lstrip())
            elif line.strip() and not re.match(r"^\s*([-*_])(?:\s*\1){2,}\s*$", line):
                paragraph = word.add_paragraph()
                self._append_inline(paragraph, line)
            index += 1
        if fenced and code_lines:
            paragraph = word.add_paragraph()
            run = paragraph.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        return referenced

    @staticmethod
    def _split_table_row(value: str) -> list[str]:
        return [cell.strip().replace("\\|", "|") for cell in value.strip().strip("|").split("|")]

    def _append_table(self, word: WordDocumentType, rows: list[list[str]]) -> None:
        columns = max((len(row) for row in rows), default=0)
        if not rows or not columns:
            return
        table = word.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        for row_index, values in enumerate(rows):
            for column_index in range(columns):
                cell = table.cell(row_index, column_index)
                cell.text = values[column_index] if column_index < len(values) else ""
                _set_cell_margins(cell)
                if row_index == 0:
                    _shade(cell, "1F4E78")
                for run in cell.paragraphs[0].runs:
                    _set_font(run, size=9.5)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)

    def _append_inline(self, paragraph: Paragraph, value: str) -> None:
        position = 0
        for match in _INLINE_TOKEN.finditer(value):
            if match.start() > position:
                run = paragraph.add_run(value[position:match.start()])
                _set_font(run)
            token = match.group(0)
            if token.startswith("**") or token.startswith("__"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
                _set_font(run)
            elif token.startswith(("*", "_")):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
                _set_font(run)
            elif token.startswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
            else:
                link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
                if link:
                    _add_hyperlink(paragraph, link.group(1), link.group(2))
            position = match.end()
        if position < len(value):
            run = paragraph.add_run(value[position:])
            _set_font(run)

    def _append_asset(
        self, word: WordDocumentType, asset: DocumentAsset, caption: str, referenced: bool = False
    ) -> None:
        paragraph = word.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        data = self.assets.bytes_for(asset)
        if asset.mime_type == "image/svg+xml":
            inline = self._add_svg(paragraph, data, asset.width, asset.height)
        else:
            width = self._display_width(asset.width, asset.height)
            inline = paragraph.add_run().add_picture(io.BytesIO(data), width=Inches(width))
        inline._inline.docPr.set("descr", f"asset-{'ref' if referenced else 'auto'}:{asset.id}")
        label = word.add_paragraph(f"图  {caption}")
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label.paragraph_format.keep_with_next = False
        caption_marker = label.add_run(_CAPTION_MARKER)
        caption_marker.font.hidden = True
        for run in label.runs:
            _set_font(run, size=9)
            run.font.italic = True

    @staticmethod
    def _display_width(width: int | None, height: int | None) -> float:
        if not width or not height:
            return 5.8
        return min(6.4, max(2.0, width / 144.0))

    def _add_svg(self, paragraph: Paragraph, data: bytes, width: int | None, height: int | None):
        display_width = self._display_width(width, height)
        ratio = (height / width) if width and height else 0.6
        placeholder = Image.new("RGBA", (1000, max(1, round(1000 * ratio))), (255, 255, 255, 0))
        buffer = io.BytesIO()
        placeholder.save(buffer, format="PNG")
        buffer.seek(0)
        inline = paragraph.add_run().add_picture(buffer, width=Inches(display_width))
        package = paragraph.part.package
        partname = package.next_partname("/word/media/image%d.svg")
        svg_part = Part(PackURI(str(partname)), "image/svg+xml", data, package)
        relationship = paragraph.part.relate_to(svg_part, RT.IMAGE)
        blip = inline._inline.graphic.graphicData.pic.blipFill.blip
        blip.set(qn("r:embed"), relationship)
        return inline

    def _append_diagrams(self, word: WordDocumentType, block: Block) -> None:
        diagrams = (block.fragment or {}).get("er_diagrams", [])
        for diagram in diagrams:
            try:
                png = self._diagram_png(diagram)
            except Exception:
                continue
            paragraph = word.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline = paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(6.35))
            inline._inline.docPr.set("descr", f"diagram:{block.id}:{diagram.get('diagram_id', 'diagram')}")
            caption = word.add_paragraph(f"图  {diagram.get('title') or 'ER 图'}")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_marker = caption.add_run(_CAPTION_MARKER)
            caption_marker.font.hidden = True
            for run in caption.runs:
                _set_font(run, size=9)
                run.font.italic = True

    def _diagram_png(self, diagram: dict[str, object]) -> bytes:
        projection = diagram_review_projection(str(diagram["drawio_xml"]))
        pages = projection.get("pages") or []
        entities = [entity for page in pages for entity in page.get("entities", [])]
        relationships = [edge for page in pages for edge in page.get("relationships", [])]
        if not entities:
            raise ValueError("diagram has no entities")
        columns = min(3, max(1, math.ceil(math.sqrt(len(entities)))))
        rows = math.ceil(len(entities) / columns)
        card_width, card_height = 460, 220
        gap_x, gap_y, margin, title_height = 90, 100, 80, 100
        canvas_width = margin * 2 + columns * card_width + (columns - 1) * gap_x
        canvas_height = title_height + margin + rows * card_height + (rows - 1) * gap_y + margin
        image = Image.new("RGB", (canvas_width, canvas_height), "white")
        draw = ImageDraw.Draw(image)
        font = self._image_font(30)
        small = self._image_font(24)
        tiny = self._image_font(20)
        title = str(diagram.get("title") or "ER 图")
        draw.text((margin, 30), title, fill="#17202A", font=font)
        boxes: dict[str, tuple[int, int, int, int]] = {}
        for index, entity in enumerate(entities):
            row, column = divmod(index, columns)
            x = margin + column * (card_width + gap_x)
            y = title_height + margin + row * (card_height + gap_y)
            boxes[str(entity.get("id"))] = (x, y, x + card_width, y + card_height)
        for edge in relationships:
            source, target = boxes.get(str(edge.get("source"))), boxes.get(str(edge.get("target")))
            if not source or not target:
                continue
            source_center = ((source[0] + source[2]) // 2, (source[1] + source[3]) // 2)
            target_center = ((target[0] + target[2]) // 2, (target[1] + target[3]) // 2)
            start = self._box_edge(source, target_center)
            end = self._box_edge(target, source_center)
            draw.line((start, end), fill="#6B7280", width=4)
            self._arrow(draw, start, end)
        for entity in entities:
            box = boxes[str(entity.get("id"))]
            draw.rounded_rectangle(box, radius=16, fill="#F8FBFF", outline="#316A9A", width=4)
            header = (box[0], box[1], box[2], box[1] + 58)
            draw.rounded_rectangle(header, radius=16, fill="#1F4E78", outline="#1F4E78")
            draw.rectangle((box[0], box[1] + 42, box[2], box[1] + 58), fill="#1F4E78")
            draw.text((box[0] + 18, box[1] + 13), str(entity.get("label") or entity.get("id")), fill="white", font=small)
            fields = [str(item.get("label") or item.get("id")) for item in entity.get("fields", [])]
            for field_index, field in enumerate(fields[:6]):
                draw.text((box[0] + 18, box[1] + 74 + field_index * 23), "• " + field[:36], fill="#263746", font=tiny)
        output = io.BytesIO()
        image.save(output, format="PNG", optimize=True)
        return output.getvalue()

    @staticmethod
    def _arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        length = 18
        points = [
            end,
            (end[0] - length * math.cos(angle - 0.45), end[1] - length * math.sin(angle - 0.45)),
            (end[0] - length * math.cos(angle + 0.45), end[1] - length * math.sin(angle + 0.45)),
        ]
        draw.polygon(points, fill="#6B7280")

    @staticmethod
    def _box_edge(box: tuple[int, int, int, int], toward: tuple[int, int]) -> tuple[int, int]:
        center_x = (box[0] + box[2]) / 2
        center_y = (box[1] + box[3]) / 2
        delta_x = toward[0] - center_x
        delta_y = toward[1] - center_y
        if not delta_x and not delta_y:
            return round(center_x), round(center_y)
        scale_x = ((box[2] - box[0]) / 2) / abs(delta_x) if delta_x else float("inf")
        scale_y = ((box[3] - box[1]) / 2) / abs(delta_y) if delta_y else float("inf")
        scale = min(scale_x, scale_y)
        return round(center_x + delta_x * scale), round(center_y + delta_y * scale)

    @staticmethod
    def _image_font(size: int):
        candidates = (
            Path("C:/Windows/Fonts/msyh.ttc"),
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        )
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
        return ImageFont.load_default()

    def _parse(
        self,
        word: WordDocumentType,
        document_id: str,
        blocks: list[Block],
        known_assets: dict[str, DocumentAsset],
    ) -> list[ImportedBlock]:
        marker_text = "\n".join(paragraph.text for paragraph in word.paragraphs[:3])
        document_marker = _DOC_MARKER.search(marker_text)
        if document_marker and document_marker.group(1) != document_id:
            raise BlockError("这个 DOCX 属于另一份 PRD", 422)

        by_id = {block.id: block for block in blocks}
        by_title: dict[str, list[Block]] = {}
        for block in blocks:
            by_title.setdefault(block.title.strip(), []).append(block)
        sections: list[dict[str, object]] = []
        current: dict[str, object] | None = None
        for item in self._body_items(word):
            if isinstance(item, Paragraph):
                marker = _BLOCK_MARKER.search(item.text)
                if marker:
                    block_id, version = marker.group(1), int(marker.group(2))
                    if block_id not in by_id:
                        raise BlockError("DOCX 包含当前文档不存在的 Block", 422)
                    title = _BLOCK_MARKER.sub("", item.text).strip() or by_id[block_id].title
                    current = {"id": block_id, "version": version, "title": title, "parts": [], "images": []}
                    sections.append(current)
                    continue
                if current is None and item.style and item.style.name.startswith("Heading"):
                    matches = by_title.get(item.text.strip(), [])
                    if len(matches) == 1:
                        block = matches[0]
                        current = {"id": block.id, "version": block.version, "title": block.title, "parts": [], "images": []}
                        sections.append(current)
                        continue
            if current is None:
                continue
            markdown, images = self._item_markdown(item, known_assets, len(current["images"]))
            if markdown:
                current["parts"].append(markdown)
            current["images"].extend(images)

        if not sections:
            raise BlockError("未识别到 firstFlight Block 标识；请使用系统导出的 DOCX 编辑后再回传", 422)
        if len({section["id"] for section in sections}) != len(sections):
            raise BlockError("DOCX 中存在重复的 Block 章节", 422)
        missing = [block.title for block in blocks if block.id not in {section["id"] for section in sections}]
        if missing:
            raise BlockError("DOCX 缺少章节：" + "、".join(missing[:10]), 422)
        return [
            ImportedBlock(
                block_id=str(section["id"]),
                base_version=int(section["version"]),
                title=str(section["title"]),
                content="\n\n".join(str(part).strip() for part in section["parts"] if str(part).strip()),
                images=list(section["images"]),
            )
            for section in sections
        ]

    @staticmethod
    def _body_items(word: WordDocumentType) -> Iterable[Paragraph | Table]:
        for child in word.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, word)
            elif child.tag == qn("w:tbl"):
                yield Table(child, word)

    def _item_markdown(
        self, item: Paragraph | Table, known_assets: dict[str, DocumentAsset], image_offset: int
    ) -> tuple[str, list[ImportedImage]]:
        if isinstance(item, Table):
            rows = [[cell.text.replace("|", "\\|").strip() for cell in row.cells] for row in item.rows]
            if not rows:
                return "", []
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            header = "| " + " | ".join(rows[0]) + " |"
            divider = "| " + " | ".join("---" for _ in range(width)) + " |"
            body = ["| " + " | ".join(row) + " |" for row in rows[1:]]
            return "\n".join((header, divider, *body)), []

        if _CAPTION_MARKER in item.text:
            return "", []

        images: list[ImportedImage] = []
        image_lines: list[str] = []
        blips = item._p.xpath(".//a:blip")
        doc_properties = item._p.xpath(".//wp:docPr")
        for index, blip in enumerate(blips):
            description = doc_properties[index].get("descr", "") if index < len(doc_properties) else ""
            if description.startswith("diagram:"):
                continue
            if description.startswith("asset-auto:") and description[11:] in known_assets:
                continue
            if description.startswith("asset-ref:") and description[10:] in known_assets:
                asset = known_assets[description[10:]]
                image_lines.append(f"![{asset.file_name}](asset://{asset.id})")
                continue
            relationship_id = blip.get(qn("r:embed"))
            part = item.part.related_parts.get(relationship_id)
            if part is None or not getattr(part, "blob", None):
                continue
            mime_type = getattr(part, "content_type", "")
            if mime_type not in {"image/png", "image/jpeg", "image/webp", "image/svg+xml"}:
                continue
            extension = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp", "image/svg+xml": ".svg"}[mime_type]
            token = f"[[FF_IMPORTED_IMAGE_{image_offset + len(images)}]]"
            images.append(ImportedImage(token, f"docx-image-{image_offset + len(images) + 1}{extension}", mime_type, part.blob))
            image_lines.append(token)

        text = self._paragraph_markdown(item)
        return "\n".join([part for part in (text, *image_lines) if part]), images

    @staticmethod
    def _paragraph_markdown(paragraph: Paragraph) -> str:
        if not paragraph.text.strip():
            return ""
        parts: list[str] = []
        for run in paragraph.runs:
            text = _DOC_MARKER.sub("", _BLOCK_MARKER.sub("", run.text))
            if not text:
                continue
            if run.bold:
                text = f"**{text}**"
            if run.italic:
                text = f"*{text}*"
            parts.append(text)
        value = "".join(parts).strip() or _DOC_MARKER.sub("", _BLOCK_MARKER.sub("", paragraph.text)).strip()
        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading"):
            try:
                level = min(6, int(style.split()[-1]))
            except ValueError:
                level = 2
            return f"{'#' * level} {value}"
        if style.startswith("List Bullet"):
            return f"- {value}"
        if style.startswith("List Number"):
            return f"1. {value}"
        return value

    @staticmethod
    def _file_name(title: str) -> str:
        safe = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", title).strip(" ._") or "PRD"
        return f"{safe[:120]}.docx"
