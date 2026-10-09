"""Extract reviewable text from a PRD supplied during project intake."""

from __future__ import annotations

import base64
import binascii
import io
import re
from collections import Counter
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from app.schemas.workflow import MAX_IMPORTED_PRD_CHARACTERS
from app.services import imported_prd_store


MAX_SOURCE_BYTES = 24 * 1024 * 1024
MAX_SOURCE_TEXT = MAX_IMPORTED_PRD_CHARACTERS
MAX_DOCX_EXPANDED_BYTES = 128 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".docx", ".md", ".pdf", ".txt"}
# Markdown has six heading levels; anything deeper has nowhere to go.
MAX_DOCX_HEADING_LEVEL = 6
_HEADING_STYLE = re.compile(r"^Heading\s+(\d+)$")
# The bare token a code block opens with — "yaml", "python", "text".
_CODE_LANGUAGE = re.compile(r"^[A-Za-z][\w+#.-]{0,20}$")


def import_prd_source(filename: str, content_base64: str, *, keep_file: bool = False) -> dict[str, str]:
    """Decode and extract a bounded, text-only baseline for downstream Agents.

    ``keep_file`` also parks the original bytes on disk and reports an
    ``import_id`` for them. Nothing in the brief grows: the id is a pointer, and
    it is what lets a later step pull the images back out of the file.
    """

    safe_name = Path(filename).name.strip()
    extension = Path(safe_name).suffix.lower()
    if not safe_name or extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("支持 DOCX、Markdown、TXT 或文字型 PDF 文件")
    if len(content_base64) > 4 * ((MAX_SOURCE_BYTES + 2) // 3):
        raise ValueError("PRD 文件不能超过 24 MB")
    try:
        raw = base64.b64decode(content_base64, validate=True)
    except (binascii.Error, ValueError) as error:
        raise ValueError("PRD 文件编码无效") from error
    if not raw:
        raise ValueError("PRD 文件不能为空")
    if len(raw) > MAX_SOURCE_BYTES:
        raise ValueError("PRD 文件不能超过 24 MB")
    if extension in {".md", ".txt"} and len(raw) > 4 * MAX_SOURCE_TEXT + 3:
        raise ValueError(f"PRD 提取文字不能超过 {MAX_SOURCE_TEXT} 个字符")

    if extension in {".md", ".txt"}:
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("文本 PRD 必须使用 UTF-8 编码") from error
    elif extension == ".pdf":
        content = _pdf_text(raw)
    else:
        content = _docx_text(raw)

    content = content.strip()
    if not content:
        raise ValueError("PRD 中没有可提取的文字")
    if len(content) > MAX_SOURCE_TEXT:
        raise ValueError(f"PRD 提取文字不能超过 {MAX_SOURCE_TEXT} 个字符")
    return {
        "filename": safe_name,
        "media_type": _media_type(extension),
        "title": _title(content, safe_name),
        "content": content,
        "import_id": imported_prd_store.save(raw) if keep_file and extension == ".docx" else "",
    }


def _pdf_text(raw: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
    except ModuleNotFoundError as error:
        raise ValueError("当前服务未安装 PDF 文字解析组件，请改用 DOCX、Markdown 或 TXT") from error
    except Exception as error:
        raise ValueError("PDF 文件无效或已损坏") from error
    if reader.is_encrypted:
        raise ValueError("不支持加密 PDF")
    if len(reader.pages) > 50:
        raise ValueError("PDF 不能超过 50 页")
    parts: list[str] = []
    total = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        total += len(text) + (2 if parts else 0)
        if total > MAX_SOURCE_TEXT:
            raise ValueError(f"PRD 提取文字不能超过 {MAX_SOURCE_TEXT} 个字符")
        parts.append(text)
    return "\n\n".join(parts)


def _docx_text(raw: bytes) -> str:
    return _docx_content(raw)[0]


def _docx_content(raw: bytes) -> tuple[str, list[dict[str, object]]]:
    """Extract ``(markdown, images)`` from a DOCX.

    Images cannot travel as text, so each one is left in the markdown as an
    ``[[FF_IMPORTED_IMAGE_n]]`` token; the caller stores the bytes and swaps the
    token for an ``asset://`` reference once the image has an id.
    """

    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 5_000 or sum(item.file_size for item in entries) > MAX_DOCX_EXPANDED_BYTES:
                raise ValueError("DOCX 解压后内容过大，请移除超大嵌入文件后重试")
    except BadZipFile as error:
        raise ValueError("DOCX 文件无效或已损坏") from error
    try:
        document = WordDocument(io.BytesIO(raw))
    except Exception as error:
        raise ValueError("DOCX 文件无效或已损坏") from error

    parts: list[str] = []
    images: list[dict[str, object]] = []
    total = 0

    def append(text: str) -> None:
        nonlocal total
        total += len(text) + (2 if parts else 0)
        if total > MAX_SOURCE_TEXT:
            raise ValueError(f"PRD 提取文字不能超过 {MAX_SOURCE_TEXT} 个字符")
        parts.append(text)

    size_levels = _docx_size_levels(document)

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text and not paragraph._p.xpath(".//a:blip"):
                continue
            style = paragraph.style.name if paragraph.style else ""
            level = _heading_level(style)
            if level is None:
                level = size_levels.get(_paragraph_font_size(paragraph))
            # A heading's runs are bold because they are a heading; keeping the
            # emphasis would put "**" inside every Block title.
            body = _paragraph_markdown(paragraph, images, emphasis=level is None)
            if not body:
                continue
            if level is not None:
                append(f"{'#' * level} {body}")
            elif style.startswith("List Bullet"):
                append(f"- {body}")
            elif style.startswith("List Number"):
                append(f"1. {body}")
            else:
                append(body)
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            single = _single_column_text(table)
            if single is not None:
                code = _code_block(single)
                if code:
                    append(code)
                elif single.strip():
                    append(single.strip())
                continue
            rows = []
            table_chars = 0
            for row in table.rows:
                # A cell holding a line break would otherwise spill out of the
                # row and break every table below it.
                cells = [re.sub(r"\s+", " ", _cell_text(cell, images))
                         .replace("|", "\\|").strip() for cell in row.cells]
                table_chars += sum(len(cell) for cell in cells)
                if total + table_chars > MAX_SOURCE_TEXT:
                    raise ValueError(f"PRD 提取文字不能超过 {MAX_SOURCE_TEXT} 个字符")
                rows.append(cells)
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            # GFM only reads a table when the header, the divider and every body
            # row sit on consecutive lines — one part, not one part per row.
            append("\n".join(
                ["| " + " | ".join(rows[0]) + " |",
                 "| " + " | ".join("---" for _ in range(width)) + " |"]
                + ["| " + " | ".join(row) + " |" for row in rows[1:]]
            ))
    return "\n\n".join(parts), images


def _runs_markdown(paragraph: Paragraph) -> str:
    """The paragraph's runs with their bold/italic/code emphasis kept."""

    parts: list[str] = []
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        if not text.strip():
            parts.append(text)
            continue
        font = run.font.name or ""
        if "Consolas" in font or font.startswith("Courier") or font.endswith("Mono"):
            text = f"`{text.strip()}`"
        if run.bold:
            text = f"**{text}**"
        elif run.italic:
            text = f"*{text}*"
        parts.append(text)
    return "".join(parts).strip() or paragraph.text.strip()


def _paragraph_markdown(
    paragraph: Paragraph, images: list[dict[str, object]], emphasis: bool = True
) -> str:
    """The paragraph's markdown, with an image token per embedded image."""

    text = _runs_markdown(paragraph) if emphasis else paragraph.text.strip()
    tokens = _image_tokens(paragraph, images)
    if not tokens:
        return text
    return "\n".join(part for part in (text, *tokens) if part)


def _cell_text(cell, images: list[dict[str, object]]) -> str:
    """Cell text, keeping emphasis and collecting any image inside it."""

    lines: list[str] = []
    for paragraph in cell.paragraphs:
        value = _paragraph_markdown(paragraph, images)
        if value.strip():
            lines.append(value.strip())
    return "\n".join(lines)


_IMAGE_EXTENSIONS = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/jpg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/webp": ".webp",
    "image/svg+xml": ".svg",
    "image/tiff": ".tiff",
    "image/emf": ".emf",
    "image/wmf": ".wmf",
}


def _image_tokens(paragraph: Paragraph, images: list[dict[str, object]]) -> list[str]:
    """One token per image in the paragraph, in document order."""

    tokens: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        relationship_id = blip.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        if not relationship_id:
            continue
        part = paragraph.part.related_parts.get(relationship_id)
        blob = getattr(part, "blob", None)
        if not blob:
            continue
        mime_type = getattr(part, "content_type", "") or "image/png"
        extension = _IMAGE_EXTENSIONS.get(mime_type, ".png")
        token = f"[[FF_IMPORTED_IMAGE_{len(images)}]]"
        images.append({
            "token": token,
            "file_name": f"image-{len(images) + 1}{extension}",
            "mime_type": mime_type,
            "content": blob,
            "width": None,
            "height": None,
        })
        tokens.append(token)
    return tokens


def extract_docx_images(raw: bytes) -> list[dict[str, object]]:
    """The images of a DOCX, in the order their tokens appear in the markdown."""

    return _docx_content(raw)[1]


def _single_column_text(table: Table) -> str | None:
    """The stacked cell text of a one-column table, else ``None``.

    A one-column table is how these PRDs carry a code sample or a callout box.
    Neither is a Markdown table, and a multi-line cell cannot be written as one:
    its inner lines escape as loose paragraphs, so a code comment starting with
    ``#`` reads as a heading and wrecks the outline.
    """

    cells = []
    for row in table.rows:
        row_cells = row.cells
        if len(row_cells) != 1:
            return None
        cells.append(row_cells[0].text)
    return "\n".join(cells) if cells else None


def _code_block(text: str) -> str | None:
    """A fenced code block when the cell opens with a bare language name."""

    lines = text.splitlines()
    if len(lines) < 2:
        return None
    language = lines[0].strip()
    if not _CODE_LANGUAGE.match(language):
        return None
    body = "\n".join(lines[1:]).strip("\n")
    if not body:
        return None
    return f"```{language}\n{body}\n```"


def _heading_level(style: str) -> int | None:
    """The level an explicit ``Heading N`` style carries, if any."""

    match = _HEADING_STYLE.match(style.strip())
    if not match:
        return None
    return min(MAX_DOCX_HEADING_LEVEL, max(1, int(match.group(1))))


def _paragraph_font_size(paragraph: Paragraph) -> float | None:
    """The largest font size used by the paragraph's visible runs."""

    sizes = [
        round(run.font.size.pt, 1)
        for run in paragraph.runs
        if run.text.strip() and run.font.size is not None
    ]
    return max(sizes) if sizes else None


def _docx_size_levels(document: WordDocument) -> dict[float, int]:
    """Heading levels for a DOCX whose outline survives only in font sizes.

    Plenty of PRD files keep every paragraph on the ``Normal`` style and mark
    chapters by font size alone — a 26pt title, 18pt chapters, 16pt sections,
    15pt and 14pt below that. Word then stores no outline at all, so without
    this the whole document reads as one flat wall of paragraphs and any
    structure an Agent invents is guesswork.

    The most common size is the body text. A size used exactly once above it is
    the document's own title and is skipped, and the sizes between rank downward
    as heading levels. A file that does use ``Heading`` styles is left to them,
    so an explicit outline always wins.
    """

    counts: Counter[float] = Counter()
    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        if _heading_level(paragraph.style.name if paragraph.style else "") is not None:
            return {}
        size = _paragraph_font_size(paragraph)
        if size is not None:
            counts[size] += 1
    if not counts:
        return {}
    body = counts.most_common(1)[0][0]
    larger = sorted((size for size in counts if size > body), reverse=True)
    if larger and counts[larger[0]] == 1:
        larger = larger[1:]
    return {
        size: level
        for level, size in enumerate(larger[:MAX_DOCX_HEADING_LEVEL], start=1)
    }


def _title(content: str, filename: str) -> str:
    for line in content.splitlines():
        candidate = re.sub(r"[*_`]+", "", line.strip().lstrip("#")).strip()
        if candidate:
            return candidate[:100]
    return Path(filename).stem[:100] or "导入的 PRD"


def _media_type(extension: str) -> str:
    return {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".md": "text/markdown",
        ".pdf": "application/pdf",
        ".txt": "text/plain",
    }[extension]
