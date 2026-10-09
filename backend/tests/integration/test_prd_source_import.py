"""Existing-PRD intake keeps the source as revision evidence."""

import base64
from collections import deque
from io import BytesIO

from fastapi.testclient import TestClient
from docx import Document
from docx.shared import Pt

from app.domain.types import ClarificationAnalysis
from app.database.models import Artifact, Project
from app.schemas.workflow import MAX_IMPORTED_PRD_CHARACTERS, ImportedPrdSource
from app.services.prd_source_import import import_prd_source
from main import create_app
from tests.helpers.factories import make_complete_brief
from tests.helpers.fake_agent import ScriptedAgentGateway


def _add_sized(document, text, points):
    run = document.add_paragraph().add_run(text)
    run.font.size = Pt(points)


def _one_cell_document(lines):
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = lines[0]
    for line in lines[1:]:
        cell.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def _encoded(value: str) -> str:
    return base64.b64encode(value.encode("utf-8")).decode("ascii")


def test_import_prd_accepts_the_character_limit_and_rejects_over_it():
    content = "# PRD\n" + "需" * (MAX_IMPORTED_PRD_CHARACTERS - len("# PRD\n"))
    imported = import_prd_source("large.md", _encoded(content))
    assert len(imported["content"]) == MAX_IMPORTED_PRD_CHARACTERS
    assert ImportedPrdSource.model_validate(imported).content == content
    try:
        import_prd_source("large.md", _encoded(content + "需"))
    except ValueError as error:
        assert str(MAX_IMPORTED_PRD_CHARACTERS) in str(error)
    else:
        raise AssertionError("oversized PRD was accepted")


def test_import_prd_extracts_markdown_for_project_creation(session_factory):
    agent = ScriptedAgentGateway()
    with TestClient(
        create_app(agent_gateway=agent, session_factory=session_factory)
    ) as client:
        response = client.post(
            "/sessions/import-prd",
            json={
                "filename": "existing-prd.md",
                "content_base64": _encoded("# Customer Portal\n\n## Current flow\nUsers sign in."),
            },
        )

    assert response.status_code == 200
    assert response.json() == {
        "filename": "existing-prd.md",
        "media_type": "text/markdown",
        "title": "Customer Portal",
        "content": "# Customer Portal\n\n## Current flow\nUsers sign in.",
        # 只有 .docx 会把原件落盘（图片回环需要原文件），md 没有 import_id。
        "import_id": "",
    }


def test_import_prd_extracts_docx_headings_and_paragraphs(session_factory):
    document = Document()
    document.add_heading("Payments Refresh", level=1)
    document.add_paragraph("Keep the existing refund workflow.")
    buffer = BytesIO()
    document.save(buffer)

    with TestClient(
        create_app(agent_gateway=ScriptedAgentGateway(), session_factory=session_factory)
    ) as client:
        response = client.post(
            "/sessions/import-prd",
            json={
                "filename": "payments-prd.docx",
                "content_base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
            },
        )

    assert response.status_code == 200
    assert response.json()["title"] == "Payments Refresh"
    assert response.json()["content"] == (
        "# Payments Refresh\n\nKeep the existing refund workflow."
    )


def test_import_prd_reads_a_docx_outline_that_lives_only_in_font_sizes():
    """Lots of real PRDs keep every paragraph on "Normal" and vary the size."""

    document = Document()
    _add_sized(document, "Payments Refresh", 26)
    _add_sized(document, "Current flow", 18)
    _add_sized(document, "Refunds", 16)
    for text in ["Refunds are manual today.", "The queue is the bottleneck.",
                 "Finance reconciles by hand.", "No retry policy exists.",
                 "Volume grows every quarter."]:
        _add_sized(document, text, 11)
    buffer = BytesIO()
    document.save(buffer)

    content = import_prd_source(
        "payments.docx", base64.b64encode(buffer.getvalue()).decode("ascii")
    )["content"]

    assert content.splitlines()[0] == "Payments Refresh"
    assert "# Current flow" in content
    assert "## Refunds" in content
    assert "Refunds are manual today." in content


def test_import_prd_fences_a_single_cell_table_that_holds_code():
    encoded = _one_cell_document([
        "yaml",
        "log_id: string          # log_ 前缀",
        "# 过程事件不进 LongTermMemory",
        "level: enum             # debug | info",
    ])

    content = import_prd_source("code.docx", encoded)["content"]

    assert content.startswith("```yaml\n")
    assert content.endswith("\n```")
    # The comment keeps its "#" instead of turning into a Markdown heading.
    assert "# 过程事件不进 LongTermMemory" in content
    assert "| yaml" not in content


def test_import_prd_keeps_a_single_cell_callout_as_plain_text():
    encoded = _one_cell_document([
        "定位：本节回答知识库与长期记忆的关系。",
        "边界：本机制不替代检索。",
    ])

    content = import_prd_source("callout.docx", encoded)["content"]

    assert content.startswith("定位：")
    assert "边界：本机制不替代检索。" in content
    assert not content.startswith("|")


def test_imported_prd_is_persisted_and_sent_to_intake_agent(session_factory):
    agent = ScriptedAgentGateway(
        analyze_results=deque(
            [ClarificationAnalysis(ready_for_spec=True, questions=[], assumptions=[])]
        )
    )
    source = {
        "filename": "baseline.md",
        "media_type": "text/markdown",
        "title": "Baseline PRD",
        "content": "# Baseline PRD\n\nThe old rule needs correction.",
    }
    brief = make_complete_brief().model_dump(mode="json")
    brief["source_prd"] = source

    with TestClient(
        create_app(agent_gateway=agent, session_factory=session_factory)
    ) as client:
        response = client.post(
            "/sessions",
            json={"request_id": "imported-prd-intake", "actor_id": "owner-1", "brief": brief},
        )

    assert response.status_code == 201
    assert agent.calls[0][0] == "analyze_brief"
    assert agent.calls[0][1]["brief"]["source_prd"] == source
    with session_factory() as db:
        project = db.query(Project).filter_by(creation_request_id="imported-prd-intake").one()
        assert project.brief["source_prd"] == source
        original = db.query(Artifact).filter_by(project_id=project.id, kind="ORIGINAL_REQUIREMENT").one()
        assert original.content["source_prd"] == source


def test_import_prd_rejects_unsupported_files(session_factory):
    with TestClient(
        create_app(agent_gateway=ScriptedAgentGateway(), session_factory=session_factory)
    ) as client:
        response = client.post(
            "/sessions/import-prd",
            json={"filename": "prd.pages", "content_base64": _encoded("content")},
        )

    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "PRD_IMPORT_INVALID"
