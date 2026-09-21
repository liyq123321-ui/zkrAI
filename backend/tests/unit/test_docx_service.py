import io
import zipfile

from PIL import Image
from docx import Document as WordDocument

from app.database.database import create_engine_for_url, init_database, make_session_factory
from app.database.models import Project
from app.models import Block, DocumentArtifact, DocumentAsset
from app.services.block_service import BlockService
from app.services.document_asset_service import DocumentAssetService
from app.services.docx_service import DocxService


ER_XML = """<mxfile compressed="false"><diagram id="page" name="Page-1"><mxGraphModel>
<root><mxCell id="0"/><mxCell id="1" parent="0"/>
<mxCell id="users" value="用户" style="shape=table;html=1;" vertex="1" parent="1"><mxGeometry x="20" y="20" width="180" height="100" as="geometry"/></mxCell>
<mxCell id="orders" value="订单" style="shape=table;html=1;" vertex="1" parent="1"><mxGeometry x="300" y="20" width="180" height="100" as="geometry"/></mxCell>
<mxCell id="rel" value="创建" edge="1" source="users" target="orders" parent="1"><mxGeometry relative="1" as="geometry"/></mxCell>
</root></mxGraphModel></diagram></mxfile>"""


def _services(tmp_path):
    engine = create_engine_for_url(f"sqlite:///{tmp_path / 'docx.db'}")
    init_database(engine)
    factory = make_session_factory(engine)
    with factory() as db:
        db.add(Project(
            id="project", session_id="session", creation_request_id="request",
            brief={"title": "共享 PRD"}, final_approver="owner", phase="SPECIFICATION",
        ))
        db.commit()
    blocks = BlockService(factory)
    assets = DocumentAssetService(blocks, tmp_path / "assets")
    return engine, blocks, assets, DocxService(blocks, assets)


def _png() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (640, 360), "#DCEEFE").save(output, format="PNG")
    return output.getvalue()


def _svg() -> bytes:
    return b'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 640 240">
    <rect width="640" height="240" rx="24" fill="#EAF4FC"/>
    <path d="M80 120h460" stroke="#1F4E78" stroke-width="8"/>
    <circle cx="80" cy="120" r="28" fill="#1F4E78"/>
    <circle cx="540" cy="120" r="28" fill="#168A42"/>
    </svg>'''


def test_docx_export_persists_artifact_and_embeds_assets_and_diagrams(tmp_path):
    engine, blocks, assets, docx = _services(tmp_path)
    try:
        snapshot = blocks.open("session", "owner")
        snapshot = blocks.import_outline(
            snapshot["document"]["id"], "owner", "custom", "# 产品背景\n# 数据模型",
            snapshot["document"]["revision"],
        )
        document_id = snapshot["document"]["id"]
        first, second = snapshot["blocks"]
        blocks.save(document_id, first["id"], "owner", first["version"], {
            "content": "原始说明\n\n| 指标 | 目标 |\n| --- | --- |\n| 可用性 | 99.9% |\n"
        })
        assets.upload(document_id, first["id"], "owner", "页面渲染.png", "image/png", _png())
        assets.upload(document_id, first["id"], "owner", "流程图.svg", "image/svg+xml", _svg())
        with blocks.write() as db:
            stored = db.get(Block, second["id"])
            stored.fragment = {"er_diagrams": [{
                "diagram_id": "data-model", "title": "核心数据模型",
                "after_section": "core_objects", "drawio_xml": ER_XML,
            }]}
        file_name, raw, revision = docx.export(document_id, "owner")
        assert file_name == "共享 PRD.docx"
        assert raw.startswith(b"PK")
        word = WordDocument(io.BytesIO(raw))
        assert "原始说明" in "\n".join(paragraph.text for paragraph in word.paragraphs)
        assert len(word.inline_shapes) == 3
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            assert any(name.endswith(".svg") for name in package.namelist())
        with blocks.factory() as db:
            artifact = db.query(DocumentArtifact).filter_by(
                document_id=document_id, revision=revision, kind="docx"
            ).one()
            assert artifact.content_hash
            assert db.query(DocumentAsset).filter_by(document_id=document_id).count() == 2
    finally:
        engine.dispose()


def test_exported_docx_round_trip_updates_the_original_block(tmp_path):
    engine, blocks, assets, docx = _services(tmp_path)
    try:
        snapshot = blocks.open("session", "owner")
        snapshot = blocks.import_outline(
            snapshot["document"]["id"], "owner", "custom", "# 产品背景",
            snapshot["document"]["revision"],
        )
        document_id = snapshot["document"]["id"]
        block = snapshot["blocks"][0]
        blocks.save(document_id, block["id"], "owner", block["version"], {"content": "原始说明\n"})
        file_name, raw, revision = docx.export(document_id, "owner")
        word = WordDocument(io.BytesIO(raw))
        target = next(paragraph for paragraph in word.paragraphs if paragraph.text == "原始说明")
        target.text = "Word 修改后的说明"
        edited = io.BytesIO()
        word.save(edited)

        result = docx.import_document(
            document_id, "owner", revision, file_name, edited.getvalue()
        )

        current = blocks.snapshot(document_id, "owner")
        updated = current["blocks"][0]
        assert result["changed_block_ids"] == [block["id"]]
        assert updated["content"] == "Word 修改后的说明\n"
        assert updated["version"] == 3
        assert len(blocks.versions(document_id, block["id"], "owner")) == 3
    finally:
        engine.dispose()


def test_docx_import_extracts_a_new_word_image_into_document_assets(tmp_path):
    engine, blocks, assets, docx = _services(tmp_path)
    try:
        snapshot = blocks.open("session", "owner")
        snapshot = blocks.import_outline(
            snapshot["document"]["id"], "owner", "custom", "# 产品背景",
            snapshot["document"]["revision"],
        )
        document_id = snapshot["document"]["id"]
        block = snapshot["blocks"][0]
        blocks.save(document_id, block["id"], "owner", block["version"], {"content": "正文\n"})
        file_name, raw, revision = docx.export(document_id, "owner")
        word = WordDocument(io.BytesIO(raw))
        word.add_paragraph().add_run().add_picture(io.BytesIO(_png()))
        edited = io.BytesIO()
        word.save(edited)

        docx.import_document(document_id, "owner", revision, file_name, edited.getvalue())

        current = blocks.snapshot(document_id, "owner")
        assert "asset://asset_" in current["blocks"][0]["content"]
        assert len(current["assets"]) == 1
        assert current["assets"][0]["metadata"] == {"source": "docx_import"}
    finally:
        engine.dispose()
