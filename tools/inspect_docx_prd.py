from __future__ import annotations

from collections import Counter
from pathlib import Path

from docx import Document


DOCX = Path(r"C:\Users\zkr\Documents\WXWork\1688857625913676\Cache\File\2026-09\人机混动软件工程平台 PRD（v0.1） (7).docx")


def main() -> None:
    document = Document(DOCX)
    print(
        "paragraphs=",
        len(document.paragraphs),
        "tables=",
        len(document.tables),
        "inline_shapes=",
        len(document.inline_shapes),
    )
    print("styles=", Counter((p.style.name if p.style else "NONE") for p in document.paragraphs))
    for index, paragraph in enumerate(document.paragraphs):
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        sizes = sorted({round(run.font.size.pt, 1) for run in paragraph.runs if run.font.size})
        bold = sum(1 for run in paragraph.runs if run.bold)
        style = paragraph.style.name if paragraph.style else "NONE"
        print(
            f"{index:03d}|{style}|sizes={sizes}|bold={bold}/{len(paragraph.runs)}|{text[:240]}"
        )


if __name__ == "__main__":
    main()
